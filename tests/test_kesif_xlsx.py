"""KESIF XLSX: kapsayici degil, sayfalardan olusan tablo dosyasi.

Onceki surum XLSX'i zip diye "kapsayici" sayip icerigine hic bakmiyordu.
Teknik olarak XLSX bir zip'tir ama isleyis olarak bir tablo dosyasidir;
sayfalar ACILABILDIGI icin bu bir olgu sorusudur, tahmin degil.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

# Kesif ekstrasi (`.[kesif]`) kurulu degilse bu modul TOPLANAMAZ: magika,
# yonlendiricinin modul seviyesinde import ettigi bir bagimlilik. Guard
# olmadan collection hatasi tum suite'i durdurur -- yalnizca kesif
# testlerini degil.
pytest.importorskip("magika", reason="kesif ekstrasi kurulu degil")

from ads.file_detection.formats import tabular  # noqa: E402
from ads.file_detection.models import Severity  # noqa: E402
from ads.file_detection.readers import read_workbook  # noqa: E402
from ads.file_detection.router import Flow, route  # noqa: E402


def _kitap(hedef: Path, bos_sayfa: bool = True, ikinci_tablo: bool = True) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "satis"
    ws.append(["urun", "adet", "tutar"])
    for r in (("klavye", 3, 450.0), ("fare", 7, 210.5), ("monitor", 2, 3400.0)):
        ws.append(list(r))
    if ikinci_tablo:
        w2 = wb.create_sheet("stok")
        w2.append(["kod", "miktar"])
        w2.append(["A1", 120])
        w2.append(["A2", 45])
    if bos_sayfa:
        wb.create_sheet("bos")
    wb.save(hedef)
    return hedef


def test_xlsx_kapsayici_degil_tablo_akisina_gidiyor(tmp_path: Path) -> None:
    k = route(_kitap(tmp_path / "kitap.xlsx"))

    assert k.format in {"xlsx", "xlsm"}
    assert k.akis is Flow.TABLE, "XLSX artik kapsayici sayilmamali"
    assert k.akis is not Flow.CONTAINER
    assert k.deterministik is True
    assert any("calisma kitabi" in a for a, _ in k.kanitlar)


def test_sayfalar_acilip_yapilari_olculuyor(tmp_path: Path) -> None:
    sonuc = tabular.inspect_structure(_kitap(tmp_path / "kitap.xlsx"))

    assert not sonuc.hata
    adlar = {s.ad for s in sonuc.sayfalar}
    assert adlar == {"satis", "stok", "bos"}
    assert {s.ad for s in sonuc.tablo_sayfalari} == {"satis", "stok"}
    bos = next(s for s in sonuc.sayfalar if s.ad == "bos")
    assert bos.tablo_mu is False


def test_birden_cok_tablo_sayfasi_tercih_sorusu_olarak_sunuluyor(tmp_path: Path) -> None:
    rapor = tabular.inspect(_kitap(tmp_path / "kitap.xlsx"))

    secimli = [b for b in rapor.bulgular if b.secenekler]
    assert len(secimli) == 1
    b = secimli[0]
    assert b.onem is Severity.WARNING
    assert "TERCIH" in b.aciklama
    sayfalar = {s.parametre["sayfa"] for s in b.secenekler}
    assert sayfalar == {"satis", "stok"}
    assert sum(1 for s in b.secenekler if s.onerilen) == 1


def test_tek_tablo_sayfasinda_secim_sorulmuyor(tmp_path: Path) -> None:
    """Tek sayfa varsa ortada bir tercih yok; gereksiz soru sorulmaz."""
    rapor = tabular.inspect(
        _kitap(tmp_path / "tek.xlsx", bos_sayfa=False, ikinci_tablo=False))

    assert not [b for b in rapor.bulgular if b.secenekler]


def test_tablo_olmayan_kitap_human_feedbacke_dusuyor(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "serbest"
    ws.append(["tek hucre"])
    p = tmp_path / "tablosuz.xlsx"
    wb.save(p)

    k = route(p)
    assert k.deterministik is False
    assert k.akis is Flow.ADJUDICATION

    rapor = tabular.inspect(p)
    kritik = [b for b in rapor.bulgular if b.onem is Severity.CRITICAL]
    assert len(kritik) == 1
    assert kritik[0].secenekler[0].parametre == {"mod": "human_feedback"}


def test_secilen_sayfa_okunuyor(tmp_path: Path) -> None:
    p = _kitap(tmp_path / "kitap.xlsx")

    varsayilan = read_workbook(p, {})
    assert varsayilan["basarili"] is True
    assert varsayilan["secim_yapildi"] is False
    assert set(varsayilan["tum_tablo_sayfalari"]) == {"satis", "stok"}

    secili = read_workbook(p, {"sayfa": "stok"})
    assert secili["sayfa"] == "stok"
    assert secili["satir_sayisi"] == 2
    assert secili["onizleme"][0]["kod"] == "A1"


def test_olmayan_sayfa_istenirse_uydurmuyor(tmp_path: Path) -> None:
    sonuc = read_workbook(_kitap(tmp_path / "kitap.xlsx"), {"sayfa": "yok"})

    assert sonuc["basarili"] is False
    assert "yok" in sonuc["hata"]
    assert "mevcut_sayfalar" in sonuc, "hangi sayfalarin oldugu soylenmeli"


def test_bozuk_kitap_cokmez(tmp_path: Path) -> None:
    p = tmp_path / "bozuk.xlsx"
    p.write_bytes(b"bu bir xlsx degil")

    sonuc = tabular.inspect_structure(p)
    assert sonuc.hata

    rapor = tabular.inspect(p)
    assert rapor.okunabilir is False


def _baslikli_kitap(hedef: Path, sayfa_adlari: tuple[str, ...]) -> Path:
    """Tablonun ustunde rapor basligi ve bos satir olan kitap.

    Kurumsal disa aktarimin olagan sekli; kenar durum degil.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for ad in sayfa_adlari:
        ws = wb.create_sheet(ad)
        ws.append(["2024 Yillik Satis Raporu"])
        ws.append([])
        ws.append(["urun_id", "adet", "tutar"])
        for i in range(5):
            ws.append([f"P{i:03d}", i, i * 10.5])
    wb.save(hedef)
    return hedef


def test_baslik_ustunde_rapor_satiri_olan_sayfa_tablo_sayiliyor(tmp_path: Path) -> None:
    """Baslik 1. satirda degilse sayfa "tablo degil" sayiliyordu.

    Olculdu: iki sayfali physicians.xlsx icin kanit "2 sayfa acildi,
    1 tanesi tablo" diyordu; ayni yanitta IKI sayfa da profillenmisti.
    `ads.intake` bu sayfayi `header_row_inferred` ile zaten yukluyor.
    """
    p = _baslikli_kitap(tmp_path / "rapor.xlsx", ("Ocak",))

    sonuc = tabular.inspect_structure(p)

    assert len(sonuc.tablo_sayfalari) == 1
    sayfa = sonuc.sayfalar[0]
    assert sayfa.tablo_mu is True
    assert sayfa.baslik_satiri == 3
    assert sayfa.basliklar[:3] == ["urun_id", "adet", "tutar"]


def test_butun_sayfalari_baslikli_kitap_human_feedbacke_dusmuyor(tmp_path: Path) -> None:
    """Asil zarar burada: butun sayfalar boyleyse kitap yargiya dusuyordu.

    Yonlendiricide `if not tablolar` dali eskalasyon uretiyordu, oysa
    dosya deterministik olarak okunabiliyor.
    """
    p = _baslikli_kitap(tmp_path / "yillik.xlsx", ("Ocak", "Subat"))

    k = route(p)

    assert k.akis is Flow.TABLE
    assert k.deterministik is True
    assert "2 tanesi tablo" in k.kanitlar[0][1]
